import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from pypdf import PdfReader

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "skills/resume/scripts"))
from exporter import export_resume
from test_assessment import make_workspace


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root, self.plan = make_workspace(self.temp.name)

    def test_pdf_docx_content_and_separate_evidence(self):
        output = export_resume(self.root, self.plan, "application", paper="a4")
        pdf = PdfReader(output / "resume.pdf")
        pdf_text = "\n".join(page.extract_text() for page in pdf.pages)
        docx_text = "\n".join(p.text for p in Document(output / "resume.docx").paragraphs)
        for text in (pdf_text, docx_text):
            self.assertIn("Alex Example", text)
            self.assertIn("Built a reporting service during an internship.", text)
            self.assertNotIn("Evidence basis", text)
            self.assertNotIn("example.invalid", text)
        self.assertAlmostEqual(float(pdf.pages[0].mediabox.width), 595.28, places=1)
        evidence = (output / "evidence.md").read_text()
        self.assertIn("../../applicants/solo/skill.md", evidence)
        manifest = json.loads((output / "manifest.json").read_text())
        for name, checksum in manifest["files"].items():
            self.assertEqual(hashlib.sha256((output / name).read_bytes()).hexdigest(), checksum)
        with self.assertRaisesRegex(ValueError, "already exists"):
            export_resume(self.root, self.plan, "application")

    def test_unicode_and_markup_are_literal(self):
        # Output wording is reviewed by the skill; here test document encoding.
        value = 'Zoë García <Research & Development> "résumé"'
        self.plan["resume"]["name"]["text"] = value
        output = export_resume(self.root, self.plan, "unicode")
        text = "\n".join(p.extract_text() for p in PdfReader(output / "resume.pdf").pages)
        self.assertIn(value, " ".join(text.split()))  # PDF layout may wrap the name.
        self.assertEqual(Document(output / "resume.docx").paragraphs[0].text, value)

    def test_unknown_glyph_fails_without_artifacts(self):
        self.plan["resume"]["name"]["text"] = "Alex \U0001f680"
        with self.assertRaisesRegex(ValueError, "font lacks"):
            export_resume(self.root, self.plan, "missing-glyph")
        self.assertEqual([p.name for p in (self.root / "outputs").iterdir()], [".gitignore"])

    def test_render_failure_cleans_staging(self):
        with patch("exporter.write_docx", side_effect=OSError("synthetic disk failure")):
            with self.assertRaisesRegex(OSError, "disk failure"):
                export_resume(self.root, self.plan, "failed")
        self.assertEqual([p.name for p in (self.root / "outputs").iterdir()], [".gitignore"])

    def test_rejects_dispute_before_export(self):
        self.plan["reviews"]["applicants/solo/skill.md"]["disputed"] = True
        with self.assertRaisesRegex(ValueError, "Disputed"):
            export_resume(self.root, self.plan, "disputed")
        self.assertFalse((self.root / "outputs/disputed").exists())

    def test_long_resume_paginates(self):
        line = self.plan["resume"]["sections"][0]["items"][0]
        self.plan["resume"]["sections"][0]["items"] = [dict(line, text=f"Entry {i}: " + line["text"]) for i in range(100)]
        output = export_resume(self.root, self.plan, "long")
        pdf = PdfReader(output / "resume.pdf")
        self.assertGreater(len(pdf.pages), 1)
        self.assertIn("Entry 99:", pdf.pages[-1].extract_text())
        self.assertIn("Entry 99:", Document(output / "resume.docx").paragraphs[-1].text)


if __name__ == "__main__":
    unittest.main()
