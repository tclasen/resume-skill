"""Produce PDF/DOCX resumes and a separate claim-to-evidence report locally."""

import hashlib
from html import escape
import json
import os
from pathlib import Path
import tempfile

from docx import Document
from docx.shared import Inches, Pt, RGBColor
import reportlab
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont, TTFError
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from assessment import evidence_report, markdown, validate_plan
from workspace import exclusive_write, identifier, load_workspace


def resume_lines(resume):
    yield "name", resume["name"]["text"]
    for line in resume["contact"]:
        yield "contact", line["text"]
    for section in resume["sections"]:
        yield "heading", section["heading"]
        for line in section["items"]:
            yield "item", line["text"]


def fonts_for(resume, font=None, bold_font=None):
    directory = Path(reportlab.__file__).resolve().parent / "fonts"
    regular_path = Path(font) if font else directory / "Vera.ttf"
    bold_path = Path(bold_font) if bold_font else (regular_path if font else directory / "VeraBd.ttf")
    fonts = []
    for path in (regular_path, bold_path):
        name = "Resume-" + hashlib.sha256(path.read_bytes()).hexdigest()[:16]
        if name not in pdfmetrics.getRegisteredFontNames():
            try:
                pdfmetrics.registerFont(TTFont(name, str(path)))
            except TTFError as error:
                raise ValueError(f"Cannot load TrueType font {path}: {error}") from error
        fonts.append(name)
    for role, text in resume_lines(resume):
        if role == "item":
            text = "\u2022 " + text
        face = pdfmetrics.getFont(fonts[1] if role in {"name", "heading"} else fonts[0]).face
        missing = sorted({ord(char) for char in text if ord(char) not in face.charToGlyph})
        if missing:
            codes = ", ".join(f"U+{value:04X}" for value in missing)
            raise ValueError(f"PDF font lacks characters {codes}; supply --font and optionally --bold-font with coverage")
    return fonts


def write_pdf(path, resume, page_size, fonts):
    regular, bold = fonts
    base = dict(fontName=regular, fontSize=10.5, leading=14, textColor=colors.HexColor("#17212b"),
                alignment=TA_LEFT, splitLongWords=True, spaceAfter=6)
    styles = {
        "name": ParagraphStyle("ResumeName", **{**base, "fontName": bold, "fontSize": 21, "leading": 25, "keepWithNext": True}),
        "contact": ParagraphStyle("ResumeContact", **{**base, "fontSize": 9.5, "leading": 12}),
        "heading": ParagraphStyle("ResumeSection", **{**base, "fontName": bold, "fontSize": 12, "spaceBefore": 12, "keepWithNext": True}),
        "item": ParagraphStyle("ResumeItem", **{**base, "leftIndent": 11, "firstLineIndent": -9}),
    }
    document = SimpleDocTemplate(str(path), pagesize=page_size, rightMargin=44, leftMargin=44,
                                 topMargin=40, bottomMargin=40, title="Resume", author="", subject="")
    story = []
    for role, text in resume_lines(resume):
        prefix = "&#8226; " if role == "item" else ""
        story.append(Paragraph(prefix + escape(text), styles[role]))
    story.append(Spacer(1, 2))
    document.build(story)


def write_docx(path, resume, page_size):
    document = Document()
    section = document.sections[0]
    section.page_width = Pt(page_size[0])
    section.page_height = Pt(page_size[1])
    section.top_margin = section.bottom_margin = Inches(0.56)
    section.left_margin = section.right_margin = Inches(0.61)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("17212B")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.title = "Resume"
    document.core_properties.subject = ""
    document.core_properties.comments = ""
    for role, text in resume_lines(resume):
        paragraph = document.add_paragraph(style="List Bullet" if role == "item" else "Normal")
        run = paragraph.add_run(text)
        if role in {"name", "heading"}:
            run.bold = True
            run.font.size = Pt(21 if role == "name" else 12)
            paragraph.paragraph_format.keep_with_next = True
            if role == "heading":
                paragraph.paragraph_format.space_before = Pt(12)
        elif role == "contact":
            run.font.size = Pt(9.5)
    document.save(path)


def export_resume(workspace, plan, name, paper="letter", font=None, bold_font=None):
    root = load_workspace(workspace)
    name = identifier(name)
    if paper not in {"letter", "a4"}:
        raise ValueError("Paper must be letter or a4")
    concepts = validate_plan(root, plan)
    resume = plan["resume"]
    if resume is None:
        raise ValueError("Plan has no resume selection; add evidence-backed resume lines before export")
    target = root / "outputs" / name
    if target.exists() or target.is_symlink():
        raise ValueError("Output id already exists; use a new id to preserve prior exports")
    fonts = fonts_for(resume, font, bold_font)
    with tempfile.TemporaryDirectory(prefix=".export-", dir=root / "outputs") as temp:
        staging = Path(temp)
        page_size = letter if paper == "letter" else A4
        write_pdf(staging / "resume.pdf", resume, page_size, fonts)
        write_docx(staging / "resume.docx", resume, page_size)
        md = []
        for role, text in resume_lines(resume):
            prefix = "# " if role == "name" else "## " if role == "heading" else "- " if role == "item" else ""
            md += [prefix + markdown(text), ""]
        exclusive_write(staging / "resume.md", "\n".join(md))
        exclusive_write(staging / "evidence.md", evidence_report(plan, concepts, link_prefix="../../"))
        manifest = {"version": 1, "applicant_snapshot": plan["applicant_snapshot"],
                    "position_snapshot": plan["position_snapshot"], "paper": paper,
                    "plan_sha256": hashlib.sha256(json.dumps(plan, sort_keys=True).encode()).hexdigest(),
                    "files": {file.name: hashlib.sha256(file.read_bytes()).hexdigest() for file in sorted(staging.iterdir())}}
        exclusive_write(staging / "manifest.json", json.dumps(manifest, indent=2) + "\n")
        # Recheck after rendering so a source edit during export fails closed.
        validate_plan(root, plan)
        target.mkdir(mode=0o700)  # Exclusive reservation; never replace another export.
        published = []
        try:
            for file in staging.iterdir():
                os.chmod(file, 0o600)
                destination = target / file.name
                file.rename(destination)
                published.append(destination)
        except Exception:
            for file in published:
                file.unlink()
            target.rmdir()
            raise
    return target
